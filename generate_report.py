"""
Generate FINAL_REPORT.docx for AI Smart Expense Tracker project.
"""
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

doc = Document()

# ── Page margins ──
for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)

style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
style.paragraph_format.line_spacing = 1.5

# ── Helper functions ──
def add_heading_styled(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0, 0, 0)
    return h

def add_para(text, bold=False, italic=False, align=None, size=12):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'justify':
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def add_bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    return p

def add_table_simple(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(11)
    return table

# ═══════════════════════════════════════════
# COVER PAGE
# ═══════════════════════════════════════════
doc.add_paragraph()
doc.add_paragraph()
add_para('AI SMART EXPENSE TRACKER', bold=True, align='center', size=24)
doc.add_paragraph()
add_para('A Project Report', align='center', size=16)
add_para('Submitted in partial fulfillment of the requirements', align='center', size=12)
add_para('for the award of the Degree of', align='center', size=12)
doc.add_paragraph()
add_para('BACHELOR OF TECHNOLOGY', bold=True, align='center', size=14)
add_para('in', align='center', size=12)
add_para('COMPUTER SCIENCE AND ENGINEERING', bold=True, align='center', size=14)
doc.add_paragraph()
doc.add_paragraph()
add_para('Submitted by:', align='center', size=12)
add_para('MAHESH SANAKA', bold=True, align='center', size=14)
doc.add_paragraph()
doc.add_paragraph()
add_para('2025 - 2026', bold=True, align='center', size=14)

doc.add_page_break()

# ═══════════════════════════════════════════
# TABLE OF CONTENTS
# ═══════════════════════════════════════════
add_para('TABLE OF CONTENTS', bold=True, align='center', size=16)
doc.add_paragraph()

toc_items = [
    ('ABSTRACT', ''),
    ('1. INTRODUCTION', ''),
    ('    1.1 Project Overview', ''),
    ('    1.2 Objective and Scope', ''),
    ('2. LITERATURE SURVEY', ''),
    ('3. SYSTEM ANALYSIS', ''),
    ('    3.1 Problem Statement', ''),
    ('    3.2 Existing System', ''),
    ('    3.3 Proposed System', ''),
    ('    3.4 Feasibility Study', ''),
    ('    3.4.1 Technical Feasibility', ''),
    ('    3.4.2 Operational Feasibility', ''),
    ('    3.4.3 Economic Feasibility', ''),
    ('4. SYSTEM SPECIFICATIONS', ''),
    ('    4.1 Hardware Specification', ''),
    ('    4.2 Software Specification', ''),
    ('    4.3 Brief Overview of Software Tools', ''),
    ('5. SYSTEM DESIGN', ''),
    ('    5.1 System Architecture', ''),
    ('    5.2 Use Case Diagram', ''),
    ('    5.3 Data Flow Diagram', ''),
    ('    5.4 ER Diagram', ''),
    ('    5.5 Modules Description', ''),
    ('6. SCREENSHOTS', ''),
    ('7. PSEUDOCODE', ''),
    ('8. TESTING AND IMPLEMENTATION', ''),
    ('9. CONCLUSION', ''),
    ('10. BIBLIOGRAPHY', ''),
]

for item, _ in toc_items:
    p = doc.add_paragraph()
    run = p.add_run(item)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    if not item.startswith('    '):
        run.bold = True
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)

doc.add_page_break()

# ═══════════════════════════════════════════
# ABSTRACT
# ═══════════════════════════════════════════
add_heading_styled('ABSTRACT', level=1)
add_para(
    'The AI Smart Expense Tracker is a full-stack web application designed to help individuals manage their personal finances intelligently. '
    'The system leverages artificial intelligence through a multi-agent architecture to provide real-time financial insights, spending analysis, '
    'budget monitoring, anomaly detection, savings coaching, and investment advisory — functioning as a personal Certified Financial Planner (CFP). '
    'Users can track daily expenses, set budgets, define financial goals, upload bulk transactions via CSV, and receive personalized recommendations '
    'based on their spending patterns.',
    align='justify'
)
add_para(
    'The application is built using React.js for the frontend with Recharts for data visualization, and Flask (Python) for the backend with SQLite '
    'as the database. The AI engine consists of six specialized agents — Spending Analyst, Budget Guardian, Anomaly Detector, Savings Coach, '
    'Investment Advisor, and Bill Tracker — each responsible for a distinct aspect of financial analysis. The system also incorporates a Financial '
    'Health Score (0–100) computed from five weighted parameters: savings rate, spending consistency, category balance, budget adherence, and '
    'emergency fund readiness.',
    align='justify'
)
add_para(
    'Key features include a comprehensive dashboard with interactive charts, the 50/30/20 budget rule analysis, monthly spending comparisons, '
    'insurance coverage tracking, investment suggestions with asset allocation, and a responsive horizontal navigation bar. The application '
    'demonstrates the practical application of AI agents in personal finance management, providing actionable insights that empower users to '
    'make informed financial decisions.',
    align='justify'
)

doc.add_page_break()

# ═══════════════════════════════════════════
# 1. INTRODUCTION
# ═══════════════════════════════════════════
add_heading_styled('1. INTRODUCTION', level=1)

add_heading_styled('1.1 Project Overview', level=2)
add_para(
    'The AI Smart Expense Tracker is a modern web-based personal finance management application that combines traditional expense tracking '
    'with artificial intelligence to deliver a comprehensive financial management experience. Unlike conventional expense trackers that merely '
    'record transactions, this system actively analyzes spending patterns, detects anomalies, monitors budgets, and provides personalized '
    'investment and savings recommendations.',
    align='justify'
)
add_para(
    'The project follows a client-server architecture with a React.js single-page application (SPA) on the frontend communicating with a '
    'Flask REST API backend. The AI capabilities are implemented through a multi-agent system where each agent specializes in a specific '
    'domain of financial analysis, collectively functioning as a virtual Certified Financial Planner.',
    align='justify'
)

add_heading_styled('1.2 Objective and Scope', level=2)
add_para('The primary objectives of the project are:', align='justify')
add_bullet('To develop an intuitive expense tracking system with categorized transaction management.')
add_bullet('To implement AI-powered agents that analyze spending patterns and provide actionable financial insights.')
add_bullet('To create a Financial Health Score that gives users a quantifiable measure of their financial wellness.')
add_bullet('To provide personalized investment suggestions, insurance recommendations, and savings strategies.')
add_bullet('To enable budget setting, financial goal tracking, and monthly spending comparisons.')
add_bullet('To support bulk transaction upload via CSV for convenience.')
add_bullet('To deliver a responsive, modern user interface with interactive data visualizations.')

add_para(
    'The scope of the project encompasses personal finance management for individual users, covering expense tracking, income management, '
    'budget monitoring, goal setting, investment advisory, insurance suggestions, and comprehensive financial health analysis through AI agents.',
    align='justify'
)

doc.add_page_break()

# ═══════════════════════════════════════════
# 2. LITERATURE SURVEY
# ═══════════════════════════════════════════
add_heading_styled('2. LITERATURE SURVEY', level=1)
add_para(
    'Personal finance management has evolved significantly with the advent of digital tools and artificial intelligence. Several studies and '
    'existing systems have influenced the design and development of this project.',
    align='justify'
)
add_para('Multi-Agent Systems in Finance', bold=True)
add_para(
    'Multi-agent systems (MAS) have been extensively used in financial applications. Research by Wooldridge (2009) demonstrates that agent-based '
    'architectures are particularly suitable for financial analysis due to the inherent modularity of financial domains — budgeting, risk '
    'assessment, investment planning, and anomaly detection can each be handled by specialized agents working collaboratively.',
    align='justify'
)
add_para('The 50/30/20 Budget Rule', bold=True)
add_para(
    'The 50/30/20 budgeting rule, popularized by Senator Elizabeth Warren in "All Your Worth: The Ultimate Lifetime Money Plan" (2005), '
    'suggests allocating 50% of after-tax income to needs, 30% to wants, and 20% to savings. This rule serves as the baseline for the '
    'budget analysis feature in our application.',
    align='justify'
)
add_para('Financial Health Scoring', bold=True)
add_para(
    'The concept of financial health scoring draws from credit scoring models like FICO and CIBIL, adapted for personal finance wellness. '
    'Our implementation uses a composite score derived from five weighted parameters, similar to approaches used by fintech companies like '
    'Mint and CRED for evaluating user financial behavior.',
    align='justify'
)
add_para('Existing Applications', bold=True)
add_para(
    'Popular expense tracking applications such as Mint, YNAB (You Need A Budget), Splitwise, and Money Manager have been studied. While '
    'these applications excel in transaction recording and basic categorization, they lack the depth of AI-powered analysis that our system '
    'provides through its multi-agent architecture. Our application differentiates itself by offering CFP-level financial advisory capabilities.',
    align='justify'
)

doc.add_page_break()

# ═══════════════════════════════════════════
# 3. SYSTEM ANALYSIS
# ═══════════════════════════════════════════
add_heading_styled('3. SYSTEM ANALYSIS', level=1)

add_heading_styled('3.1 Problem Statement', level=2)
add_para(
    'Managing personal finances effectively is a challenge for many individuals. Most people lack access to professional financial planning '
    'services due to high costs. Existing expense tracking applications provide basic recording functionality but fail to deliver intelligent, '
    'actionable insights. There is a need for an affordable, AI-powered solution that not only tracks expenses but also acts as a personal '
    'financial advisor — analyzing spending patterns, detecting unusual expenditures, recommending budget adjustments, suggesting investments, '
    'and monitoring financial health in real-time.',
    align='justify'
)

add_heading_styled('3.2 Existing System', level=2)
add_para(
    'Current expense tracking solutions have the following limitations:',
    align='justify'
)
add_bullet('Manual categorization with no intelligent analysis of spending patterns.')
add_bullet('No multi-agent AI system for comprehensive financial advice.')
add_bullet('Lack of Financial Health Score to quantify overall financial wellness.')
add_bullet('No personalized investment or insurance recommendations based on user profiles.')
add_bullet('Limited monthly comparison and trend analysis capabilities.')
add_bullet('No integration of the 50/30/20 budget rule for automated assessment.')

add_heading_styled('3.3 Proposed System', level=2)
add_para(
    'The proposed AI Smart Expense Tracker addresses these limitations by incorporating:',
    align='justify'
)
add_bullet('Six specialized AI agents providing CFP-level financial insights.')
add_bullet('A composite Financial Health Score (0-100) based on five weighted parameters.')
add_bullet('Automated 50/30/20 budget rule analysis with visual comparison.')
add_bullet('Personalized investment suggestions with recommended asset allocation.')
add_bullet('Insurance coverage tracking and recommendations based on user profiles.')
add_bullet('Interactive dashboard with monthly comparisons and trend analysis.')
add_bullet('Bulk CSV upload for convenient transaction management.')
add_bullet('Real-time anomaly detection for unusual spending patterns.')

add_heading_styled('3.4 Feasibility Study', level=2)

add_heading_styled('3.4.1 Technical Feasibility', level=3)
add_para(
    'The project is technically feasible as it utilizes well-established and widely-supported technologies. React.js provides a robust '
    'frontend framework with an extensive ecosystem. Flask offers a lightweight yet powerful backend framework for Python. SQLite provides '
    'a serverless database suitable for single-user applications. The AI agents are implemented using Python with Pandas and NumPy for '
    'data analysis, eliminating the need for expensive external AI APIs.',
    align='justify'
)

add_heading_styled('3.4.2 Operational Feasibility', level=3)
add_para(
    'The application features an intuitive user interface with a modern horizontal navigation bar, interactive charts, and clear visual '
    'indicators. Users can easily add transactions, set budgets, define goals, and understand their financial health through the dashboard. '
    'The system requires minimal training and is designed for users with varying levels of technical expertise.',
    align='justify'
)

add_heading_styled('3.4.3 Economic Feasibility', level=3)
add_para(
    'The project uses entirely open-source technologies — React.js, Flask, SQLite, Recharts, and Python data science libraries — resulting '
    'in zero licensing costs. The AI agents run locally without requiring paid API subscriptions. The application can be deployed on minimal '
    'hardware, making it economically viable for individual users and small-scale deployment.',
    align='justify'
)

doc.add_page_break()

# ═══════════════════════════════════════════
# 4. SYSTEM SPECIFICATIONS
# ═══════════════════════════════════════════
add_heading_styled('4. SYSTEM SPECIFICATIONS', level=1)

add_heading_styled('4.1 Hardware Specification', level=2)
add_table_simple(
    ['Component', 'Minimum Requirement'],
    [
        ['Processor', 'Intel Core i3 or equivalent'],
        ['RAM', '4 GB (8 GB recommended)'],
        ['Storage', '500 MB free disk space'],
        ['Display', '1280 x 720 resolution or higher'],
        ['Network', 'Internet connection for initial setup'],
    ]
)
doc.add_paragraph()

add_heading_styled('4.2 Software Specification', level=2)
add_table_simple(
    ['Software', 'Version'],
    [
        ['Operating System', 'Windows 10/11, macOS, or Linux'],
        ['Node.js', 'v18.0 or higher'],
        ['Python', '3.10 or higher'],
        ['Web Browser', 'Chrome, Firefox, Edge (latest)'],
        ['npm', '9.0 or higher'],
        ['pip', '23.0 or higher'],
    ]
)
doc.add_paragraph()

add_heading_styled('4.3 Brief Overview of Software Tools', level=2)

add_para('Frontend Technologies', bold=True)
add_table_simple(
    ['Technology', 'Version', 'Purpose'],
    [
        ['React.js', '18.2.0', 'Frontend UI library for building single-page applications'],
        ['React Router DOM', '6.21.0', 'Client-side routing and navigation'],
        ['Recharts', '2.10.3', 'Data visualization library for charts and graphs'],
        ['React Icons', '4.12.0', 'Icon library (Feather Icons, Tabler Icons)'],
        ['React Toastify', '9.1.3', 'Toast notification system'],
        ['Axios', '1.6.2', 'HTTP client for API communication'],
    ]
)
doc.add_paragraph()

add_para('Backend Technologies', bold=True)
add_table_simple(
    ['Technology', 'Version', 'Purpose'],
    [
        ['Flask', '3.0.0', 'Python web framework for REST API'],
        ['Flask-CORS', '4.0.0', 'Cross-Origin Resource Sharing support'],
        ['Pandas', '2.1.4', 'Data analysis and manipulation library'],
        ['NumPy', '1.26.2', 'Numerical computing library'],
        ['PyJWT', '2.8.0', 'JSON Web Token authentication'],
        ['SQLite3', 'Built-in', 'Embedded relational database'],
        ['Werkzeug', '3.0.1', 'Password hashing and WSGI utilities'],
    ]
)

doc.add_page_break()

# ═══════════════════════════════════════════
# 5. SYSTEM DESIGN
# ═══════════════════════════════════════════
add_heading_styled('5. SYSTEM DESIGN', level=1)

add_heading_styled('5.1 System Architecture', level=2)
add_para(
    'The AI Smart Expense Tracker follows a three-tier client-server architecture:',
    align='justify'
)
add_para('Presentation Layer (Frontend)', bold=True)
add_para(
    'React.js single-page application with component-based architecture. The frontend communicates with the backend via RESTful API calls '
    'using Axios. Key pages include Dashboard, Add Transaction, Transaction History, AI Agents Hub, Financial Goals, Investment Suggestions, '
    'and Upload CSV. Authentication state is managed through React Context API.',
    align='justify'
)
add_para('Business Logic Layer (Backend)', bold=True)
add_para(
    'Flask REST API with JWT-based authentication. The backend processes all business logic including transaction management, income tracking, '
    'budget management, and AI agent orchestration. Six specialized AI agents are coordinated by an Orchestrator that runs all agents in '
    'sequence and aggregates their insights.',
    align='justify'
)
add_para('Data Layer (Database)', bold=True)
add_para(
    'SQLite relational database with six tables: users, transactions, additional_incomes, budgets, financial_goals, and recurring_bills. '
    'The schema supports referential integrity through foreign key constraints.',
    align='justify'
)

doc.add_paragraph()
add_para('AI Multi-Agent Architecture', bold=True, size=13)
add_table_simple(
    ['Agent', 'Responsibility'],
    [
        ['Spending Analyst', 'Analyzes spending patterns, calculates Financial Health Score, implements 50/30/20 rule analysis, provides month-over-month comparisons'],
        ['Budget Guardian', 'Monitors budget adherence, calculates daily safe spending rates, alerts when budgets are exceeded or at risk'],
        ['Anomaly Detector', 'Detects unusual spending spikes using statistical analysis (1.5x threshold), identifies week-over-week changes'],
        ['Savings Coach', 'Tracks savings velocity, calculates monthly requirements per goal, suggests auto-savings amounts, warns about at-risk goals'],
        ['Investment Advisor', 'Provides CFP-level advice including emergency fund analysis, debt-to-income ratio, tax-saving suggestions (80C), age-based asset allocation'],
        ['Bill Tracker', 'Identifies recurring expenses, tracks bills as percentage of income, detects newly recurring expenses, predicts upcoming bills'],
    ]
)

doc.add_page_break()

add_heading_styled('5.2 Use Case Diagram', level=2)
add_para('The system supports the following use cases for the primary actor (User):', align='justify')
add_para('')
add_table_simple(
    ['Use Case', 'Description'],
    [
        ['Register / Login', 'User creates an account or authenticates with existing credentials'],
        ['Add Transaction', 'User records a new expense with date, amount, category, payment mode, and description'],
        ['Upload CSV', 'User bulk-uploads transactions from a CSV file'],
        ['View Dashboard', 'User views financial summary, health score, charts, and AI insights'],
        ['Set Budget', 'User defines monthly spending limits per category'],
        ['Create Financial Goal', 'User sets savings goals with target amounts and deadlines'],
        ['Add Funds to Goal', 'User logs contributions toward a financial goal'],
        ['View AI Insights', 'User reviews recommendations from all six AI agents'],
        ['View Investment Suggestions', 'User gets personalized investment and insurance recommendations'],
        ['Compare Months', 'User selects two months for detailed spending comparison'],
        ['Update Profile', 'User provides age, dependents, and insurance information'],
        ['Manage Income', 'User sets monthly salary and logs additional income sources'],
    ]
)

doc.add_page_break()

add_heading_styled('5.3 Data Flow Diagram', level=2)
add_para('Level 0 — Context Diagram', bold=True)
add_para(
    'The user interacts with the AI Smart Expense Tracker system by providing transaction data, income information, budget limits, and '
    'profile details. The system processes this data through its AI agents and returns financial insights, health scores, charts, budget '
    'alerts, investment suggestions, and insurance recommendations.',
    align='justify'
)
add_para('')
add_para('Level 1 — DFD', bold=True)
add_para(
    'The Level 1 DFD decomposes the system into the following processes:',
    align='justify'
)
add_bullet('P1: Authentication — Handles user registration and login, generates JWT tokens.')
add_bullet('P2: Transaction Management — Processes individual and bulk (CSV) transaction entries.')
add_bullet('P3: Dashboard Engine — Aggregates data for financial summary, health score, and 50/30/20 analysis.')
add_bullet('P4: AI Agent Orchestrator — Coordinates six AI agents to generate financial insights.')
add_bullet('P5: Budget Manager — Tracks budget limits and monitors spending against limits.')
add_bullet('P6: Goal Tracker — Manages financial goals, tracks contributions, and calculates progress.')
add_bullet('P7: Investment Engine — Analyzes profile and spending to generate investment and insurance suggestions.')

add_heading_styled('5.4 ER Diagram', level=2)
add_para('The Entity-Relationship model consists of the following entities and relationships:', align='justify')
add_para('')
add_table_simple(
    ['Entity', 'Key Attributes', 'Relationships'],
    [
        ['Users', 'id, name, email, password_hash, monthly_income, age, dependents, insurance status', 'Has many Transactions, Budgets, Goals, Incomes'],
        ['Transactions', 'id, user_id, date, amount, category, payment_mode, description', 'Belongs to User'],
        ['Additional Incomes', 'id, user_id, date, amount, source, description', 'Belongs to User'],
        ['Budgets', 'id, user_id, category, monthly_limit', 'Belongs to User (unique per category)'],
        ['Financial Goals', 'id, user_id, name, target_amount, current_amount, deadline, status', 'Belongs to User'],
        ['Recurring Bills', 'id, user_id, name, estimated_amount, category, frequency', 'Belongs to User'],
    ]
)

doc.add_page_break()

add_heading_styled('5.5 Modules Description', level=2)

modules = [
    ('Authentication Module', 'Handles user registration with password hashing (Werkzeug), login with JWT token generation, and token-based route protection. Supports session management through React Context API on the frontend.'),
    ('Dashboard Module', 'The central module displaying financial overview including income banner, Financial Health Score (0-100) with five-component breakdown, 50/30/20 budget rule analysis, CFP insights, stat cards (monthly expenses, transactions, savings, remaining balance), insurance coverage status, interactive charts (line, bar, pie), and monthly comparison with selectable months.'),
    ('Transaction Management Module', 'Allows users to add individual transactions with date, amount, category (8 categories), payment mode (5 modes), and description. Also supports bulk upload via CSV with column mapping and validation.'),
    ('AI Agents Module', 'Houses the six specialized AI agents coordinated by the Orchestrator. Each agent analyzes transaction data through different lenses and produces categorized insights with severity levels (critical, warning, success, info). The Agents Hub page displays all insights organized by agent.'),
    ('Budget Module', 'Enables users to set monthly spending limits per category. The Budget Guardian agent monitors adherence and provides daily safe spending recommendations. Budget status is reflected in the Financial Health Score.'),
    ('Financial Goals Module', 'Users can create savings goals with target amounts, deadlines, icons, and colors. The system tracks contributions, calculates progress percentages, and the Savings Coach agent provides velocity analysis and goal-at-risk warnings.'),
    ('Investment Suggestions Module', 'Provides personalized investment recommendations with recommended asset allocation (pie chart), spending tips, and quick financial facts. Also includes insurance suggestions (term life and health) with expandable plan details, premium estimates, and provider recommendations.'),
    ('Income Management Module', 'Manages monthly salary settings and additional income tracking. Users can log extra income from various sources (freelance, bonus, rental, etc.) with the total reflected across all financial calculations.'),
]

for name, desc in modules:
    add_para(name, bold=True)
    add_para(desc, align='justify')

doc.add_page_break()

# ═══════════════════════════════════════════
# 6. SCREENSHOTS
# ═══════════════════════════════════════════
add_heading_styled('6. SCREENSHOTS', level=1)
add_para(
    'The following screenshots demonstrate the key interfaces of the AI Smart Expense Tracker application:',
    align='justify'
)
doc.add_paragraph()

screenshots = [
    'Screenshot 1: Login Page — User authentication interface with email and password fields.',
    'Screenshot 2: Registration Page — New user signup with name, email, and password.',
    'Screenshot 3: Dashboard (Top) — Income banner, Financial Health Score (91/100), 50/30/20 rule, and CFP Insights.',
    'Screenshot 4: Dashboard (Stats) — Monthly expense cards, savings till date, health and term insurance coverage status.',
    'Screenshot 5: Dashboard (Charts) — Monthly Expense Trend line chart and Category Breakdown donut chart with date range.',
    'Screenshot 6: Dashboard (Charts) — Spending by Category bar chart and Payment Methods donut chart.',
    'Screenshot 7: Dashboard (Comparison) — Monthly Comparison with 12-month overview bar chart and selectable month dropdowns.',
    'Screenshot 8: Add Transaction Page — Form with date, amount, category, payment mode, and description fields.',
    'Screenshot 9: Transaction History Page — Tabular list of all transactions with filtering and search.',
    'Screenshot 10: AI Agents Hub — Six agent cards with real-time insights and severity indicators.',
    'Screenshot 11: Financial Goals Page — Goal cards with progress bars, contribution tracking, and deadline status.',
    'Screenshot 12: Investment Suggestions — Personalized recommendations, asset allocation pie chart, and insurance plans.',
    'Screenshot 13: Upload CSV Page — Drag-and-drop CSV upload with column mapping interface.',
]

for s in screenshots:
    add_para(s, italic=True, align='justify')
    doc.add_paragraph()  # placeholder space for screenshot

add_para('[Note: Screenshots to be captured from the running application and inserted here.]', italic=True, align='center')

doc.add_page_break()

# ═══════════════════════════════════════════
# 7. PSEUDOCODE
# ═══════════════════════════════════════════
add_heading_styled('7. PSEUDOCODE', level=1)

add_para('7.1 Financial Health Score Calculation', bold=True)
p = doc.add_paragraph()
code = """FUNCTION calculate_health_score(user_data, transactions):
    total_income = user.monthly_income
    avg_monthly_expense = MEAN(monthly_totals)

    // Savings Score (max 30 points)
    savings_rate = (total_income - avg_monthly_expense) / total_income * 100
    savings_score = MIN(30, MAX(0, savings_rate * 0.6))

    // Consistency Score (max 20 points)
    volatility = STD(monthly_totals) / MEAN(monthly_totals)
    consistency_score = MAX(0, 20 - (volatility * 20))

    // Balance Score (max 15 points)
    max_category_pct = MAX(category_percentages)
    balance_score = MAX(0, 15 - MAX(0, (max_category_pct - 30) * 0.3))

    // Budget Adherence Score (max 15 points)
    budgets_met = COUNT(budgets WHERE spending <= limit)
    adherence_score = (budgets_met / total_budgets) * 15

    // Emergency Readiness Score (max 20 points)
    monthly_savings = total_income - avg_monthly_expense
    months_saved = (monthly_savings * num_months) / avg_monthly_expense
    emergency_score = MIN(20, months_saved * 3.3)

    health_score = MIN(100, SUM(all_scores))
    RETURN health_score, breakdown
END FUNCTION"""
run = p.add_run(code)
run.font.name = 'Courier New'
run.font.size = Pt(9)
doc.add_paragraph()

add_para('7.2 AI Agent Orchestrator', bold=True)
p = doc.add_paragraph()
code2 = """FUNCTION run_all_agents(user_id):
    context = load_user_context(user_id)
    agents = [SpendingAnalyst, BudgetGuardian, AnomalyDetector,
              SavingsCoach, InvestmentAdvisor, BillTracker]

    all_insights = []
    FOR EACH agent IN agents:
        insights = agent.analyze(context)
        all_insights.APPEND(insights)
    END FOR

    SORT all_insights BY severity (critical > warning > info > success)
    RETURN all_insights
END FUNCTION"""
run = p.add_run(code2)
run.font.name = 'Courier New'
run.font.size = Pt(9)
doc.add_paragraph()

add_para('7.3 Anomaly Detection', bold=True)
p = doc.add_paragraph()
code3 = """FUNCTION detect_anomalies(transactions):
    THRESHOLD = 1.5
    category_averages = GROUP_BY(transactions, category).MEAN()
    current_month = FILTER(transactions, current_month)

    anomalies = []
    FOR EACH category IN current_month:
        current_spending = SUM(category.amounts)
        average_spending = category_averages[category]
        IF current_spending > average_spending * THRESHOLD:
            anomalies.APPEND({
                category: category,
                excess: current_spending - average_spending,
                severity: "warning"
            })
        END IF
    END FOR

    // Week-over-week comparison
    this_week_total = SUM(last_7_days)
    last_week_total = SUM(7_to_14_days_ago)
    IF this_week_total > last_week_total * THRESHOLD:
        anomalies.APPEND(weekly_spike_alert)
    END IF

    RETURN anomalies
END FUNCTION"""
run = p.add_run(code3)
run.font.name = 'Courier New'
run.font.size = Pt(9)
doc.add_paragraph()

add_para('7.4 50/30/20 Budget Rule Analysis', bold=True)
p = doc.add_paragraph()
code4 = """FUNCTION analyze_50_30_20(transactions, income):
    NEEDS_CATEGORIES = {Food, Bills, Healthcare, Transport, Education}
    WANTS_CATEGORIES = {Shopping, Entertainment, Other}

    needs_total = SUM(transactions WHERE category IN NEEDS)
    wants_total = SUM(transactions WHERE category IN WANTS)
    income_basis = income * num_months

    needs_pct = (needs_total / income_basis) * 100
    wants_pct = (wants_total / income_basis) * 100
    savings_pct = 100 - needs_pct - wants_pct

    // Compare with ideal: 50% Needs, 30% Wants, 20% Savings
    RETURN {needs_pct, wants_pct, savings_pct}
END FUNCTION"""
run = p.add_run(code4)
run.font.name = 'Courier New'
run.font.size = Pt(9)

doc.add_page_break()

# ═══════════════════════════════════════════
# 8. TESTING AND IMPLEMENTATION
# ═══════════════════════════════════════════
add_heading_styled('8. TESTING AND IMPLEMENTATION', level=1)

add_para('8.1 Testing Strategy', bold=True)
add_para(
    'The application was tested using a combination of manual testing, API testing, and visual verification. The following testing approaches were employed:',
    align='justify'
)

add_para('Unit Testing', bold=True)
add_bullet('Backend API endpoints tested individually using HTTP requests.')
add_bullet('AI agent outputs verified against known transaction datasets.')
add_bullet('Health score calculation validated with synthetic data producing expected scores (35 → 91).')
add_bullet('Database CRUD operations tested for all six tables.')

add_para('Integration Testing', bold=True)
add_bullet('Frontend-backend communication tested across all API endpoints.')
add_bullet('JWT authentication flow tested for registration, login, and protected routes.')
add_bullet('Agent orchestrator tested to ensure all six agents execute without errors.')
add_bullet('CSV upload tested with various file formats and edge cases.')

add_para('UI/UX Testing', bold=True)
add_bullet('Responsive design verified across desktop (1280px+), tablet (768px), and mobile (< 600px) viewports.')
add_bullet('Chart rendering verified for overlapping labels — resolved with inside-slice labels for pie charts.')
add_bullet('Navigation tested for all routes with active state highlighting.')
add_bullet('Form validation tested for required fields and data types.')

add_para('')
add_para('8.2 Test Results', bold=True)
add_table_simple(
    ['Test Case', 'Expected Result', 'Actual Result', 'Status'],
    [
        ['User Registration', 'Account created with hashed password', 'Account created successfully', 'PASS'],
        ['User Login', 'JWT token returned', 'Token generated and stored', 'PASS'],
        ['Add Transaction', 'Transaction saved to database', 'Transaction recorded correctly', 'PASS'],
        ['CSV Upload', 'Bulk transactions imported', 'All rows imported with validation', 'PASS'],
        ['Health Score Calc', 'Score 80+ for healthy spending', 'Score: 91/100 (Excellent)', 'PASS'],
        ['50/30/20 Analysis', 'Correct percentage split', 'Needs: 44.9%, Wants: 19.2%, Savings: 35.9%', 'PASS'],
        ['AI Agent Insights', 'Insights generated for all agents', 'All 6 agents returned insights', 'PASS'],
        ['Monthly Comparison', 'Compare any 2 months', 'Month selectors working correctly', 'PASS'],
        ['Budget Monitoring', 'Alert when budget exceeded', 'Budget Guardian alerts generated', 'PASS'],
        ['Responsive Layout', 'Proper rendering on mobile', 'Hamburger menu and stacked layout working', 'PASS'],
    ]
)

add_para('')
add_para('8.3 Implementation', bold=True)
add_para(
    'The application is implemented as a local web application with the following deployment steps:',
    align='justify'
)
add_bullet('Backend: Flask server running on port 5000 with SQLite database in /data directory.')
add_bullet('Frontend: React development server running on port 3000 with proxy to backend.')
add_bullet('Database initialization is automatic on first server startup via init_db().')
add_bullet('Authentication uses JWT tokens stored in localStorage with 24-hour expiry.')
add_bullet('AI agents are triggered on-demand when the user navigates to Dashboard or Agents Hub pages.')

doc.add_page_break()

# ═══════════════════════════════════════════
# 9. CONCLUSION
# ═══════════════════════════════════════════
add_heading_styled('9. CONCLUSION', level=1)
add_para(
    'The AI Smart Expense Tracker successfully demonstrates the application of artificial intelligence in personal finance management. '
    'By implementing a multi-agent architecture with six specialized AI agents, the system provides comprehensive financial analysis that '
    'goes far beyond traditional expense tracking applications.',
    align='justify'
)
add_para(
    'The Financial Health Score feature provides users with an intuitive, quantifiable measure of their financial wellness, while the '
    '50/30/20 budget rule analysis offers a proven framework for evaluating spending habits. The monthly comparison feature enables users '
    'to track their progress over time, and the investment suggestions module provides actionable recommendations tailored to individual '
    'profiles.',
    align='justify'
)
add_para(
    'Key achievements of the project include:',
    align='justify'
)
add_bullet('Successfully implemented a multi-agent AI system with six specialized financial analysis agents.')
add_bullet('Achieved a comprehensive Financial Health Score algorithm with five weighted parameters.')
add_bullet('Built an interactive, responsive dashboard with real-time data visualization.')
add_bullet('Developed personalized investment and insurance recommendation engines.')
add_bullet('Created a scalable architecture that can be extended with additional agents and features.')

add_para('Future Enhancements', bold=True)
add_bullet('Integration with banking APIs for automatic transaction import.')
add_bullet('Machine learning models for spending prediction and budget optimization.')
add_bullet('Mobile application using React Native for cross-platform support.')
add_bullet('Multi-currency support for international users.')
add_bullet('Advanced NLP for natural language transaction entry (e.g., "Spent 500 on lunch").')
add_bullet('Social features for shared expense tracking and group goals.')

doc.add_page_break()

# ═══════════════════════════════════════════
# 10. BIBLIOGRAPHY
# ═══════════════════════════════════════════
add_heading_styled('10. BIBLIOGRAPHY', level=1)

refs = [
    'React.js Documentation — https://react.dev/ — Official React documentation for building user interfaces.',
    'Flask Documentation — https://flask.palletsprojects.com/ — Official Flask web framework documentation.',
    'Recharts Library — https://recharts.org/ — React charting library documentation.',
    'Pandas Documentation — https://pandas.pydata.org/docs/ — Python data analysis library.',
    'SQLite Documentation — https://www.sqlite.org/docs.html — SQLite embedded database documentation.',
    'Warren, E. & Tyagi, A.W. (2005). "All Your Worth: The Ultimate Lifetime Money Plan." Free Press.',
    'Wooldridge, M. (2009). "An Introduction to MultiAgent Systems." John Wiley & Sons.',
    'JWT.io — https://jwt.io/ — JSON Web Token specification and documentation.',
    'React Router Documentation — https://reactrouter.com/ — Client-side routing for React applications.',
    'MDN Web Docs — https://developer.mozilla.org/ — Web development reference and tutorials.',
]

for i, ref in enumerate(refs, 1):
    p = doc.add_paragraph()
    run = p.add_run(f'[{i}] {ref}')
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)

# ── Save ──
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'FINAL_REPORT.docx')
doc.save(output_path)
print(f'Report saved to: {output_path}')
